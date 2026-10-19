"""Word, Excel, Markdown and CSV uploads (suveryn_rag.formats): passages with their section or cells,
nothing left out, and sūveryn's export mark found wherever the export puts it."""
from datetime import datetime
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
from suveryn_rag import formats
from suveryn_rag.formats import (
    Passage,
    cell_text,
    column_letter,
    kind_of,
    read,
    rows_in_passages,
    sheet_ref,
    split_text,
)
from suveryn_rag.marks import find_mark

FIXTURES = Path(__file__).with_name("fixtures")
LIMIT = 1152  # chunk_max_tokens * 3, as the reader is configured


def texts(reading):
    return [(p.headings, p.text, p.place) for p in reading.passages]


# ---------------------------------------------------------------------- the pieces
def test_kinds_by_extension():
    assert [kind_of(n) for n in ("a.PDF", "b.docx", "c.xlsx", "d.md", "e.markdown", "f.csv")] == \
        ["pdf", "docx", "xlsx", "md", "md", "csv"]
    assert [kind_of(n) for n in ("a.doc", "b.xls", "c.xlsm", "d.txt", "noextension")] == [None] * 5


def test_cell_references_are_written_as_excel_writes_them():
    assert [column_letter(n) for n in (1, 26, 27, 52, 703)] == ["A", "Z", "AA", "AZ", "AAA"]
    assert sheet_ref("Prices") == "Prices" and sheet_ref("Sheet_2") == "Sheet_2"
    assert sheet_ref("Price list") == "'Price list'" and sheet_ref("O'Brien") == "'O''Brien'"
    assert sheet_ref("2026") == "'2026'"


def test_cell_values_as_text():
    assert [cell_text(v) for v in (None, 300000.0, 0.21, True, "  two\nlines ")] == ["", "300000", "0.21", "TRUE", "two lines"]
    assert cell_text(datetime(2026, 10, 19)) == "2026-10-19"  # noqa: DTZ001 - Excel dates have no time zone
    assert cell_text(datetime(2026, 10, 19, 14, 30)) == "2026-10-19 14:30"  # noqa: DTZ001


def test_long_text_is_cut_at_a_sentence_or_line_and_nothing_is_lost():
    text = " ".join(f"Sentence number {i} says something." for i in range(100))
    pieces = split_text(text, 200)
    assert all(len(p) <= 200 for p in pieces) and all(p.endswith(".") for p in pieces)
    assert " ".join(pieces) == text
    assert split_text("x" * 450, 200) == ["x" * 200, "x" * 200, "x" * 50]  # no space at all: a hard cut


def test_rows_are_split_with_the_first_row_repeated_and_their_cells_named():
    rows = [(1, ["Item", "Amount"])] + [(n, [f"item {n}", str(n * 100)]) for n in range(2, 40)]
    passages = rows_in_passages(rows, 200, sheet="Prices", cells=True)
    assert len(passages) > 1 and passages[0].place.startswith("Prices!A1:B")
    for p in passages[1:]:
        assert p.text.startswith("| Item | Amount |\n")
        assert len(p.text) <= 200
    first_rows = [int(p.place.split("!A")[1].split(":")[0]) for p in passages]
    last_rows = [int(p.place.rsplit("B", 1)[1]) for p in passages]
    assert first_rows[0] == 1 and last_rows[-1] == 39
    assert all(first == last + 1 for first, last in zip(first_rows[1:], last_rows, strict=False))  # no row twice, none left out


# ---------------------------------------------------------------------- Word
def word(tmp_path, build) -> Path:
    d = docx.Document()
    build(d)
    path = tmp_path / "test.docx"
    d.save(path)
    return path


def test_word_passages_carry_their_section_and_tables_their_rows(tmp_path):
    def build(d):
        d.add_paragraph("Deed of sale", style="Title")
        d.add_paragraph("Before us, the notary.")
        d.add_heading("Article 1. Parties", 1)
        d.add_paragraph("The seller is A. The buyer is B.")
        d.add_heading("Article 2", 1)
        d.add_heading("2.1 Price", 2)
        d.add_paragraph("The price is EUR 300,000.")
        t = d.add_table(rows=3, cols=2)
        for (r, c), v in {(0, 0): "Item", (0, 1): "Amount", (1, 0): "Price", (1, 1): "300,000"}.items():
            t.cell(r, c).text = v
        t.cell(2, 0).merge(t.cell(2, 1)).text = "Paid on signing"
        d.add_heading("2.2 Costs", 2)  # nothing under it: still a passage
        d.add_heading("Article 3. Occupancy", 1)
        d.add_paragraph("Vacant.")

    r = read(word(tmp_path, build), "docx", LIMIT)
    assert texts(r) == [
        (["Deed of sale"], "Before us, the notary.", None),
        (["Deed of sale", "Article 1. Parties"], "The seller is A. The buyer is B.", None),
        (["Deed of sale", "Article 2", "2.1 Price"], "The price is EUR 300,000.", None),
        (["Deed of sale", "Article 2", "2.1 Price"], "| Item | Amount |\n| Price | 300,000 |\n| Paid on signing |", None),
        (["Deed of sale", "Article 2"], "2.2 Costs", None),
        (["Deed of sale", "Article 3. Occupancy"], "Vacant.", None),
    ]
    assert r.furniture == []


def test_word_text_in_content_controls_and_tracked_changes(tmp_path):
    from docx.oxml import parse_xml

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

    def build(d):
        body = d.element.body
        sect = body[-1]
        body.insert(len(body) - 1, parse_xml(
            f'<w:sdt {w}><w:sdtContent><w:p><w:r><w:t>Filled in a content control.</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        body.insert(len(body) - 1, parse_xml(
            f'<w:p {w}><w:r><w:t xml:space="preserve">Kept </w:t></w:r><w:del w:id="1" w:author="x"><w:r><w:delText>deleted</w:delText></w:r></w:del>'
            f'<w:ins w:id="2" w:author="x"><w:r><w:t>inserted</w:t></w:r></w:ins></w:p>'))
        assert sect is body[-1]

    r = read(word(tmp_path, build), "docx", LIMIT)
    assert [p.text for p in r.passages] == ["Filled in a content control.\n\nKept inserted"]  # one section: one passage


def test_a_paragraph_with_an_outline_level_is_a_heading(tmp_path):
    from docx.oxml import parse_xml

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

    def build(d):
        body = d.element.body
        body.insert(0, parse_xml(f'<w:p {w}><w:pPr><w:outlineLvl w:val="0"/></w:pPr><w:r><w:t>Conditions</w:t></w:r></w:p>'))
        d.add_paragraph("No suspensive conditions.")

    assert texts(read(word(tmp_path, build), "docx", LIMIT)) == [(["Conditions"], "No suspensive conditions.", None)]


def test_word_footers_are_furniture_where_the_mark_is_found(tmp_path):
    def build(d):
        d.add_paragraph("The answer.")
        d.sections[0].footer.paragraphs[0].text = "Made by sūveryn: answer · sūveryn-ref A-0123abcd"

    r = read(word(tmp_path, build), "docx", LIMIT)
    assert [p.text for p in r.passages] == ["The answer."]  # the footer is not a passage
    assert find_mark(r.furniture)["ref"] == "A-0123abcd"


def test_a_word_answer_exported_by_the_chat_carries_the_mark():
    """fixtures/marked-answer.docx: an answer exported to Word by the chat UI (lib/office.ts), as
    downloaded; the mark is its footer line. The same file saved as PDF is marked-answer.pdf."""
    r = read(FIXTURES / "marked-answer.docx", "docx", LIMIT)
    assert r.passages and r.passages[0].headings == ["What is the price?"]
    found = find_mark([*r.furniture, *(p.text for p in r.passages)])
    assert found and found["ref"] == "A-e063906e" and found["kind"] == "answer"


# ---------------------------------------------------------------------- Excel and CSV
def test_every_sheet_by_rows_with_their_cells(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["Item", "Amount", "Date"])
    ws.append(["Price", 300000, datetime(2026, 10, 19)])  # noqa: DTZ001 - Excel dates have no time zone
    ws.append([])                     # an empty row is skipped, its number isn't reused
    ws.append([None, None, "late"])   # a row starting in column C
    hidden = wb.create_sheet("Price list")
    hidden.sheet_state = "hidden"     # hidden sheets are read too: their text is in the file
    hidden["B2"] = "=1+1"             # a formula never written by Excel has no value to read
    hidden["A3"] = "Note"
    path = tmp_path / "t.xlsx"
    wb.save(path)
    assert texts(read(path, "xlsx", LIMIT)) == [
        ([], "| Item | Amount | Date |\n| Price | 300000 | 2026-10-19 |\n|  |  | late |", "Prices!A1:C4"),
        ([], "| Note |", "'Price list'!A3:A3"),
    ]


def test_an_excel_table_exported_by_the_chat_carries_the_mark():
    """fixtures/marked-table.xlsx: a table exported to Excel by the chat UI (lib/office.ts); the mark is
    the last row of the sources sheet."""
    r = read(FIXTURES / "marked-table.xlsx", "xlsx", LIMIT)
    assert [p.place for p in r.passages] == ["Table!A1:C2", "Sources!A1:D4"]
    found = find_mark([p.text for p in r.passages])
    assert found and found["ref"] == "A-31710045"


def test_csv_with_semicolons_commas_or_tabs_and_in_windows_encoding(tmp_path):
    for content, encoding in [("Name;Amount\nA;1,5\nRené;2\n", "cp1252"), ("Name,Amount\nA,\"1,5\"\nRené,2\n", "utf-8-sig"),
                              ("Name\tAmount\nA\t1,5\nRené\t2\n", "utf-8")]:
        path = tmp_path / "t.csv"
        path.write_bytes(content.encode(encoding))
        assert texts(read(path, "csv", LIMIT)) == [([], "| Name | Amount |\n| A | 1,5 |\n| René | 2 |", "A1:B3")]


def test_a_spreadsheet_too_large_is_refused(tmp_path, monkeypatch):
    monkeypatch.setattr(formats, "MAX_CELLS", 5)
    path = tmp_path / "t.csv"
    path.write_text("a;b\nc;d\ne;f\n")
    with pytest.raises(ValueError, match="more than 5 filled cells"):
        read(path, "csv", LIMIT)


# ---------------------------------------------------------------------- Markdown
def test_markdown_by_section_with_comments_as_furniture(tmp_path):
    path = tmp_path / "t.md"
    path.write_text("# Deed\n\nIntroduction.\n\n## Price\nThe price is EUR 300,000.\n\nOccupancy\n---------\n\nVacant.\n\n"
                    "# Annexes\n<!-- Made by sūveryn: answer · sūveryn-ref A-0123abcd -->\n"
                    "```\n# a comment in code, not a heading\n```\n\n# Nothing under this\n", encoding="utf-8")
    r = read(path, "md", LIMIT)
    assert texts(r) == [
        (["Deed"], "Introduction.", None),
        (["Deed", "Price"], "The price is EUR 300,000.", None),
        (["Deed", "Occupancy"], "Vacant.", None),
        (["Annexes"], "```\n# a comment in code, not a heading\n```", None),
        ([], "Nothing under this", None),
    ]
    assert find_mark(r.furniture)["ref"] == "A-0123abcd"


# ---------------------------------------------------------------------- broken files
@pytest.mark.parametrize("kind, content, says", [
    ("docx", b"PK\x03\x04 not really a zip", "not a valid Word file"),
    ("xlsx", b"PK\x03\x04 not really a zip", "not a valid Excel file"),
])
def test_a_broken_office_file_is_refused_with_a_reason(tmp_path, kind, content, says):
    path = tmp_path / f"t.{kind}"
    path.write_bytes(content)
    with pytest.raises(ValueError, match=says):
        read(path, kind, LIMIT)


def test_a_word_file_that_is_really_a_workbook_or_unpacks_too_large_is_refused(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    path = tmp_path / "t.docx"
    wb.save(path)
    with pytest.raises(ValueError, match="not a valid Word file"):
        read(path, "docx", LIMIT)
    monkeypatch.setattr(formats, "MAX_UNPACKED_BYTES", 100)
    with pytest.raises(ValueError, match="too large once unpacked"):
        read(path, "xlsx", LIMIT)


def test_passages_are_plain_data():
    assert Passage("x") == Passage("x", [], None)


def test_the_reader_reads_these_formats_without_the_pdf_pipeline(tmp_path):
    """The reader process's Extractor sends a .docx, .xlsx, .md or .csv to formats (no OCR, no layout model)."""
    from suveryn_rag.config import RagSettings
    from suveryn_rag.extract import Extractor

    reader = Extractor.__new__(Extractor)  # without loading Docling's models
    reader.settings = RagSettings(chunk_max_tokens=100)
    path = tmp_path / "notes.md"
    path.write_text("# Price\n\n" + "The price is fixed. " * 40)
    ex = reader.extract(path)
    assert ex.document is None and ex.pages == 0 and ex.furniture == []
    assert all(p.headings == ["Price"] and len(p.text) <= 300 for p in ex.passages) and len(ex.passages) > 1
