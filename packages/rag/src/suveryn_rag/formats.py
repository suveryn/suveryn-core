"""Reading Word (.docx), Excel (.xlsx), Markdown (.md) and CSV (.csv) uploads.

PDFs go through OCR and Docling's layout model (``extract``). These formats carry their text and
structure in the file itself, so they are read directly, in the same sandboxed reader process
(``isolation``): an Office file is a zip of XML, untrusted input like any PDF.

What comes out is a list of passages, each with where it sits:

- **Word and Markdown**: the section, as the chain of headings above the passage ("Article 2 ›
  The price"). A Word heading is a paragraph in a Heading or Title style, or one with an outline
  level. Tables are kept as rows of cells; a long table is split by rows, its first row repeated.
- **Excel and CSV**: the cells, as a range in Excel's own notation ("Prices!A2:D31", "A2:D31" in a
  CSV): language-neutral, and what a spreadsheet user types to find them. Values are read as
  last calculated (never formulas). The first non-empty row of a sheet is repeated at the top
  of every later passage of that sheet, so each passage says what its columns are.

Nothing is left out of the body: a heading with no text under it becomes a passage of its own.
Text outside the body (Word headers and footers, Markdown comments) is returned separately as
``furniture``: it is checked for sūveryn's export mark (``marks``) but not stored as passages,
as a PDF's repeating headers and footers aren't either. Tracked insertions are read, deletions
not: the document as it now reads.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from .reading import Unreadable

# File types by extension. The upload check (``looks_like``) and the reader confirm the content.
KINDS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".md": "md", ".markdown": "md", ".csv": "csv"}
SUFFIX = {"pdf": ".pdf", "docx": ".docx", "xlsx": ".xlsx", "md": ".md", "csv": ".csv"}

# Limits for one file, checked before and while reading, so one upload can't exhaust the reader.
MAX_UNPACKED_BYTES = 300 * 1024 * 1024   # an Office file's content once unzipped
MAX_ZIP_ENTRIES = 10_000
MAX_CELLS = 250_000                      # non-empty cells in a workbook or CSV
MAX_CHARS = 5_000_000                    # text in one document (a 60-page deed is ~150,000)


@dataclass(frozen=True)
class Passage:
    """A stored passage: its text, the headings above it, and for a spreadsheet the cell range."""

    text: str
    headings: list[str] = field(default_factory=list)
    place: str | None = None


@dataclass
class Reading:
    """What was read from one file: the passages in reading order, and the text outside the body."""

    passages: list[Passage]
    furniture: list[str] = field(default_factory=list)


def kind_of(filename: str) -> str | None:
    """"pdf", "docx", "xlsx", "md" or "csv" by the file's extension; None for anything else."""
    return KINDS.get(Path(filename).suffix.lower())


def looks_like(kind: str, head: bytes) -> bool:
    """Whether the file's first bytes fit its kind: a PDF header, a zip for Office files, and no NUL
    byte in text files (a binary file renamed to .csv or .md)."""
    if kind == "pdf":
        return head.startswith(b"%PDF-")
    if kind in ("docx", "xlsx"):
        return head.startswith(b"PK\x03\x04")
    return b"\x00" not in head


def read(path: Path, kind: str, limit: int) -> Reading:
    """Read one non-PDF file of ``kind``; passages of at most ``limit`` characters where they can be split.
    Raises ValueError with a message that can be shown when the file can't be read."""
    reader = {"docx": read_docx, "xlsx": read_xlsx, "md": read_markdown, "csv": read_csv}.get(kind)
    if reader is None:
        raise Unreadable(f"no reader for {kind} files", "unsupported")
    reading = reader(Path(path), limit)
    if sum(len(p.text) for p in reading.passages) > MAX_CHARS:
        raise Unreadable("the document has too much text to be read", "too_large")
    return reading


# ---------------------------------------------------------------------- packing text into passages
def split_text(text: str, limit: int) -> list[str]:
    """``text`` in pieces of at most ``limit`` characters, cut at a line end or a sentence end when possible."""
    pieces = []
    text = text.strip()
    while len(text) > limit:
        window = text[:limit]
        cut = max(window.rfind("\n"), window.rfind(". "), window.rfind("; "))
        cut = cut + 1 if cut > limit // 3 else (window.rfind(" ") if window.rfind(" ") > limit // 3 else limit)
        pieces.append(text[:cut].strip())
        text = text[cut:].strip()
    return [*pieces, text] if text else pieces


def pack(items: list[tuple[list[str], str]], limit: int) -> list[Passage]:
    """Consecutive pieces of text under the same headings joined into passages of up to ``limit`` characters."""
    out: list[Passage] = []
    current: list[str] = []
    headings: list[str] | None = None

    def flush() -> None:
        if current:
            out.append(Passage("\n\n".join(current), list(headings or [])))
            current.clear()

    for heads, text in items:
        for piece in split_text(text, limit):
            if heads != headings or (current and sum(len(c) + 2 for c in current) + len(piece) > limit):
                flush()
                headings = heads
            current.append(piece)
    flush()
    return out


def row_line(cells: list[str]) -> str:
    """One table row as a line: the cells between pipes, as in a Markdown table."""
    return "| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |"


def rows_in_passages(rows: list[tuple[int, list[str]]], limit: int, headings: list[str] | None = None,
                     sheet: str | None = None, cells: bool = False) -> list[Passage]:
    """Rows (number, cells) as passages of up to ``limit`` characters, the first row repeated at the top
    of every later passage. With ``cells``, each passage's place is its cell range (``sheet``! A1:D9)."""
    if not rows:
        return []
    header = row_line(rows[0][1])
    out: list[Passage] = []
    group: list[tuple[int, list[str]]] = []

    def flush() -> None:
        if not group:
            return
        lines = [row_line(c) for _, c in group]
        if group[0][0] != rows[0][0]:
            lines.insert(0, header)
        place = None
        if cells:
            last_col = max(len(c) for _, c in group)
            place = (f"{sheet_ref(sheet)}!" if sheet is not None else "") + f"A{group[0][0]}:{column_letter(last_col)}{group[-1][0]}"
        out.append(Passage("\n".join(lines), list(headings or []), place))
        group.clear()

    size = 0
    for number, values in rows:
        line = len(row_line(values)) + 1
        if group and size + line > limit:
            flush()
            size = len(header) + 1
        group.append((number, values))
        size += line
    flush()
    return out


def column_letter(n: int) -> str:
    """Excel's column name for column ``n`` (1 = A, 27 = AA)."""
    letters = ""
    n = max(n, 1)
    while n:
        n, r = divmod(n - 1, 26)
        letters = chr(65 + r) + letters
    return letters


def sheet_ref(name: str) -> str:
    """A sheet name as Excel writes it in a reference: quoted when it holds anything but letters, digits and _."""
    return name if re.fullmatch(r"[^\W\d]\w*", name) else "'" + name.replace("'", "''") + "'"


def cell_text(value) -> str:
    """A cell's value as text: whole numbers without ".0", dates as ISO dates, line breaks as spaces."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return " ".join(str(value).split())


def trimmed(cells: list[str]) -> list[str]:
    """The row without its empty cells at the end."""
    while cells and not cells[-1]:
        cells = cells[:-1]
    return cells


# ---------------------------------------------------------------------- Word
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)


def check_zip(path: Path, member: str, what: str) -> None:
    """An Office file is a zip; refuse one that isn't, lacks its main part, or unpacks to too much."""
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
    except zipfile.BadZipFile:
        raise Unreadable(f"the file is not a valid {what} file", "damaged") from None
    if member not in {i.filename for i in infos}:
        raise Unreadable(f"the file is not a valid {what} file", "damaged")
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES:
        raise Unreadable(f"the {what} file is too large once unpacked", "too_large")


def _text_of(element) -> str:
    """The visible text inside a Word element: text runs (also in hyperlinks, fields' results, content
    controls and text boxes), tabs and breaks; deleted text and field codes are other elements."""
    parts = []
    for el in element.iter(f"{W}t", f"{W}tab", f"{W}br", f"{W}cr"):
        parts.append(el.text or "" if el.tag == f"{W}t" else ("\t" if el.tag == f"{W}tab" else "\n"))
    return "".join(parts).strip()


def _heading_level(paragraph, styles: dict[str, int]) -> int | None:
    """0 (Title) to 9 for a heading paragraph, None for body text."""
    ppr = paragraph.find(f"{W}pPr")
    if ppr is not None:
        outline = ppr.find(f"{W}outlineLvl")
        if outline is not None and (v := outline.get(f"{W}val", "")).isdigit() and int(v) < 9:
            return int(v) + 1
        style = ppr.find(f"{W}pStyle")
        if style is not None:
            return styles.get(style.get(f"{W}val", ""))
    return None


def _heading_styles(document) -> dict[str, int]:
    """Style id -> heading level: the built-in Title (0) and Heading 1-9 styles (Word stores their names in
    English whatever the interface language), and any style with an outline level."""
    levels: dict[str, int] = {}
    for style in document.styles.element.iter(f"{W}style"):
        sid = style.get(f"{W}styleId", "")
        name = style.find(f"{W}name")
        name = (name.get(f"{W}val", "") if name is not None else "").strip()
        outline = style.find(f"{W}pPr/{W}outlineLvl")
        if name.lower() == "title":  # above every heading: Heading 1 sections sit under it
            levels[sid] = 0
        elif (m := HEADING_STYLE.match(name)) and m.group(1) != "0":
            levels[sid] = int(m.group(1))
        elif outline is not None and (v := outline.get(f"{W}val", "")).isdigit() and int(v) < 9:
            levels[sid] = int(v) + 1
    return levels


def _table_rows(table) -> list[list[str]]:
    """A Word table's rows as cell texts; a cell merged across columns counts once."""
    rows = []
    for tr in table.iter(f"{W}tr"):
        if tr.getparent() is not table:
            continue  # a row of a table nested in a cell: part of that cell's text
        cells = [_text_of(tc).replace("\n", " ") for tc in tr.findall(f"{W}tc")]
        if any(cells):
            rows.append(trimmed(cells))
    return rows


def read_docx(path: Path, limit: int) -> Reading:
    """A Word document's body by section, its tables by rows; headers and footers as furniture."""
    import docx

    check_zip(path, "word/document.xml", "Word")
    try:
        document = docx.Document(str(path))
    except Exception:  # noqa: BLE001 - python-docx raises many kinds for a broken file
        raise Unreadable("the file is not a valid Word file", "damaged") from None
    styles = _heading_styles(document)
    stack: list[list] = []  # [level, heading, has text under it]
    items: list[tuple[str, list[str], object]] = []  # ("text", headings, text) or ("table", headings, rows)

    def headings() -> list[str]:
        return [h for _, h, _ in stack]

    def close_to(level: int) -> None:
        while stack and stack[-1][0] >= level:
            _, text, used = stack.pop()
            if not used:  # a heading with nothing under it is text too
                items.append(("text", headings(), text))

    def add(kind: str, content) -> None:
        for entry in stack:
            entry[2] = True
        items.append((kind, headings(), content))

    def walk(container) -> None:
        for el in container:
            if el.tag == f"{W}p":
                if not (text := _text_of(el)):
                    continue
                level = _heading_level(el, styles)
                if level is None:
                    add("text", text)
                else:
                    close_to(level)
                    stack.append([level, " ".join(text.split()), False])
            elif el.tag == f"{W}tbl":
                if rows := _table_rows(el):
                    add("table", rows)
            elif el.tag == f"{W}sdt":  # a content control around paragraphs or tables
                if (content := el.find(f"{W}sdtContent")) is not None:
                    walk(content)

    walk(document.element.body)
    close_to(0)

    passages: list[Passage] = []
    texts: list[tuple[list[str], str]] = []
    for kind, heads, content in items:
        if kind == "text":
            texts.append((heads, content))
        else:
            passages += pack(texts, limit) + rows_in_passages(list(enumerate(content, 1)), limit, heads)
            texts = []
    passages += pack(texts, limit)

    furniture = []
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            try:
                if not part.is_linked_to_previous and (text := _text_of(part._element)):
                    furniture.append(text)
            except Exception:  # noqa: BLE001, S112 - a header or footer that can't be read only loses the mark check
                continue
    return Reading(passages, list(dict.fromkeys(furniture)))


# ---------------------------------------------------------------------- Excel
def read_xlsx(path: Path, limit: int) -> Reading:
    """Every worksheet's non-empty rows, as cell ranges; hidden sheets too (their text is in the file)."""
    import openpyxl

    check_zip(path, "xl/workbook.xml", "Excel")
    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises many kinds for a broken file
        raise Unreadable("the file is not a valid Excel file", "damaged") from None
    passages: list[Passage] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()  # some writers store wrong dimensions; read the rows as they are
            rows = []
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                row = trimmed([cell_text(v) for v in values])
                if row:
                    cells += sum(1 for c in row if c)
                    if cells > MAX_CELLS:
                        raise Unreadable(f"the workbook has more than {MAX_CELLS:,} filled cells".replace(",", " "), "too_large")
                    rows.append((number, row))
            passages += rows_in_passages(rows, limit, sheet=sheet.title, cells=True)
    finally:
        workbook.close()
    return Reading(passages)


# ---------------------------------------------------------------------- text files
def decode(data: bytes) -> str:
    """A text file's content: UTF-8 (with or without a byte-order mark), else Windows-1252, the code page
    Excel and older tools use for "CSV" and text in Belgium."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


class Semicolons(csv.excel):
    delimiter = ";"


def read_csv(path: Path, limit: int) -> Reading:
    """A CSV file's rows, as cell ranges. Semicolons, commas or tabs, whichever the file uses."""
    text = decode(path.read_bytes())
    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:  # one column, or too irregular to tell: semicolons, as Belgian Excel writes them
        dialect = Semicolons if sample.count(";") >= sample.count(",") else csv.excel
    rows = []
    cells = 0
    for number, values in enumerate(csv.reader(io.StringIO(text), dialect), 1):
        row = trimmed([" ".join(v.split()) for v in values])
        if row:
            cells += sum(1 for c in row if c)
            if cells > MAX_CELLS:
                raise Unreadable(f"the file has more than {MAX_CELLS:,} filled cells".replace(",", " "), "too_large")
            rows.append((number, row))
    return Reading(rows_in_passages(rows, limit, cells=True))


ATX = re.compile(r"^ {0,3}(#{1,6})[ \t]+(.*?)(?:[ \t]+#+)?[ \t]*$")
SETEXT = re.compile(r"^ {0,3}(=+|-+)[ \t]*$")
FENCE = re.compile(r"^ {0,3}(```|~~~)")
COMMENT = re.compile(r"<!--(.*?)-->", re.DOTALL)


def read_markdown(path: Path, limit: int) -> Reading:
    """A Markdown file by section (# headings and underlined ones); comments as furniture."""
    text = decode(path.read_bytes()).replace("\r\n", "\n")
    furniture = [" ".join(c.split()) for c in COMMENT.findall(text) if c.strip()]
    text = COMMENT.sub("", text)
    stack: list[list] = []  # [level, heading, has text under it]
    items: list[tuple[list[str], str]] = []
    paragraph: list[str] = []
    in_fence = False

    def headings() -> list[str]:
        return [h for _, h, _ in stack]

    def flush() -> None:
        body = "\n".join(paragraph).strip()
        paragraph.clear()
        if body:
            for entry in stack:
                entry[2] = True
            items.append((headings(), body))

    def close_to(level: int) -> None:
        flush()
        while stack and stack[-1][0] >= level:
            _, text, used = stack.pop()
            if not used:  # a heading with nothing under it is text too
                items.append((headings(), text))

    for line in text.split("\n"):
        if FENCE.match(line):
            in_fence = not in_fence
            paragraph.append(line)
        elif in_fence:
            paragraph.append(line)
        elif m := ATX.match(line):
            close_to(len(m.group(1)))
            stack.append([len(m.group(1)), m.group(2).strip(), False])
        elif (m := SETEXT.match(line)) and len(paragraph) == 1 and paragraph[0].strip():
            title = paragraph.pop().strip()
            level = 1 if m.group(1)[0] == "=" else 2
            close_to(level)
            stack.append([level, title, False])
        elif not line.strip():
            flush()
        else:
            paragraph.append(line)
    close_to(0)
    return Reading(pack(items, limit), furniture)
